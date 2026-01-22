# =========================================================
# INSIGHTIA — Analyse des verbatims clients
# Version "SaaS" : storytelling + sections + livrables
# =========================================================

import os
import io
import csv
import re
import hashlib
from datetime import datetime

import numpy as np
import pandas as pd
import streamlit as st
import matplotlib.pyplot as plt

from sklearn.cluster import KMeans
from sklearn.feature_extraction.text import TfidfVectorizer

# PDF (optionnel)
try:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    from reportlab.lib.units import cm
    REPORTLAB_AVAILABLE = True
except Exception:
    REPORTLAB_AVAILABLE = False

# PPTX (optionnel)
try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    PPTX_AVAILABLE = True
except Exception:
    PPTX_AVAILABLE = False


# =========================
# CONFIG
# =========================
APP_NAME = "INSIGHTIA"
TAGLINE = "Transformer les verbatims clients en thèmes clairs et priorités actionnables"
LOGO_PATH = os.path.join("assets", "logo_insightia.png")

st.set_page_config(page_title=APP_NAME, page_icon="📊", layout="wide")


# =========================
# STYLE (SaaS premium / sobre)
# =========================
st.markdown("""
<style>
  :root{
    --bg:#F6F8FC;
    --card:#FFFFFF;
    --stroke:#E6EAF2;
    --text:#0F172A;
    --muted:#64748B;

    --brand1:#1D4ED8;
    --brand2:#0F2E6B;

    --shadow:0 14px 35px rgba(15,23,42,0.10);
    --shadowSoft:0 10px 18px rgba(15,23,42,0.06);

    --rLG:22px;
    --rMD:14px;
  }

  .stApp{ background: var(--bg); }
  div.block-container{
    max-width: 1180px;
    padding-top: 1.0rem;
    padding-bottom: 2.2rem;
  }

  /* Sidebar */
  section[data-testid="stSidebar"]{
    background:#FFFFFF;
    border-right:1px solid var(--stroke);
  }
  section[data-testid="stSidebar"] > div{
    padding-top: 1rem;
  }

  /* Top bar spacing */
  header[data-testid="stHeader"]{
    background: transparent;
  }

  /* HERO */
  .hero{
    background: linear-gradient(90deg, var(--brand1), var(--brand2));
    border-radius: 28px;
    padding: 22px 22px;
    color:#fff;
    border: 1px solid rgba(255,255,255,0.18);
    box-shadow: var(--shadow);
    margin-bottom: 14px;
  }
  .hero-badge{
    display:inline-block;
    padding: 6px 10px;
    border-radius: 999px;
    background: rgba(255,255,255,0.16);
    border: 1px solid rgba(255,255,255,0.22);
    font-weight: 900;
    font-size: 12px;
  }
  .hero-title{
    margin-top: 10px;
    font-size: 44px;
    font-weight: 950;
    letter-spacing: -0.03em;
    line-height: 1.05;
  }
  .hero-sub{
    margin-top: 8px;
    font-size: 15px;
    opacity: 0.95;
    max-width: 900px;
  }

  /* CARDS */
  .card{
    background: var(--card);
    border: 1px solid var(--stroke);
    border-radius: var(--rLG);
    padding: 16px 18px;
    box-shadow: var(--shadowSoft);
    margin-bottom: 14px;
  }
  .card-title{
    font-size: 18px;
    font-weight: 950;
    letter-spacing: -0.02em;
    color: var(--text);
    margin-bottom: 6px;
  }
  .muted{ color: var(--muted); }

  /* Stepper */
  .stepper{
    display:flex; gap:10px; flex-wrap:wrap;
    margin-top: 10px;
  }
  .step{
    display:inline-flex;
    align-items:center;
    gap:10px;
    padding:10px 12px;
    border-radius: 16px;
    border:1px solid var(--stroke);
    background:#F8FAFF;
    font-weight: 900;
    font-size: 13px;
    color:#1E3A8A;
  }
  .step small{
    font-weight:700;
    color: var(--muted);
  }

  /* Tabs */
  button[data-baseweb="tab"]{
    font-weight: 900 !important;
    border-radius: 14px !important;
    padding: 10px 14px !important;
  }
  button[data-baseweb="tab"][aria-selected="true"]{
    background: #EEF2FF !important;
    border: 1px solid #C7D2FE !important;
  }

  /* DataFrame */
  div[data-testid="stDataFrame"]{
    border-radius: 14px;
    overflow: hidden;
    border: 1px solid var(--stroke);
  }

  /* Footer */
  .footer{
    margin-top: 56px;
    padding: 28px 24px;
    background:#FFFFFF;
    border-top: 1px solid var(--stroke);
    border-radius: var(--rLG);
    box-shadow: 0 -8px 24px rgba(15,23,42,0.05);
  }
</style>
""", unsafe_allow_html=True)


# =========================
# UI helpers
# =========================
def render_hero():
    col_logo, col_txt = st.columns([1, 8])
    with col_logo:
        if os.path.exists(LOGO_PATH):
            st.image(LOGO_PATH, width=78)
        else:
            st.markdown("### INSIGHTIA")
    with col_txt:
        st.markdown(f"""
        <div class="hero">
          <div class="hero-badge">Customer Insights</div>
          <div class="hero-title">{APP_NAME}</div>
          <div class="hero-sub">{TAGLINE}</div>
          <div class="stepper">
            <div class="step">1. Import <small>CSV & colonne texte</small></div>
            <div class="step">2. Analyse <small>thèmes dominants</small></div>
            <div class="step">3. Priorisation <small>backlog exploitable</small></div>
            <div class="step">4. Livrables <small>CSV / PDF / PPTX</small></div>
          </div>
        </div>
        """, unsafe_allow_html=True)


def card(title: str, subtitle: str = ""):
    st.markdown("<div class='card'>", unsafe_allow_html=True)
    st.markdown(f"<div class='card-title'>{title}</div>", unsafe_allow_html=True)
    if subtitle:
        st.markdown(f"<div class='muted' style='margin-top:-2px; margin-bottom:10px;'>{subtitle}</div>", unsafe_allow_html=True)


def end_card():
    st.markdown("</div>", unsafe_allow_html=True)


def fmt_int(n):
    try:
        return f"{int(n):,}".replace(",", " ")
    except Exception:
        return str(n)


# =========================
# Data / NLP helpers
# =========================
def detect_sep(file) -> str:
    try:
        sample = file.read(4096).decode("utf-8", errors="ignore")
        file.seek(0)
        return csv.Sniffer().sniff(sample).delimiter
    except Exception:
        try:
            file.seek(0)
        except Exception:
            pass
        return ";"


def read_csv_safely(file, sep: str) -> pd.DataFrame:
    return pd.read_csv(file, sep=sep, on_bad_lines="skip")


def clean_text_fast(texts) -> list[str]:
    out = []
    for t in texts:
        t = "" if t is None else str(t)
        t = t.lower()
        t = re.sub(r"\d+", " ", t)
        t = re.sub(r"[^\w\sàâçéèêëîïôûùüÿñæœ-]", " ", t)
        t = re.sub(r"\s+", " ", t).strip()
        out.append(t)
    return out


def build_cache_key(file_bytes: bytes, sep: str, text_col: str, max_rows: int, nr_topics: int) -> str:
    h = hashlib.sha256(file_bytes).hexdigest()[:16]
    return f"{h}|{sep}|{text_col}|{max_rows}|{nr_topics}"


def compute_topic_labels_from_X(X, topic_np: np.ndarray, feature_names: np.ndarray, top_n_words: int = 3) -> dict:
    # SAFE: indexation scipy avec np.where (pas Series)
    labels = {}
    for t in np.unique(topic_np):
        rows = np.where(topic_np == t)[0]
        if len(rows) == 0:
            labels[int(t)] = f"Topic {int(t)}"
            continue
        mean_tfidf = X[rows].mean(axis=0).A1
        top_terms = feature_names[mean_tfidf.argsort()[-top_n_words:]][::-1]
        labels[int(t)] = " / ".join(top_terms.tolist())
    return labels


def plot_top_topics(series: pd.Series, title: str, top_n: int = 15):
    top = series.value_counts().head(top_n).sort_values()
    fig = plt.figure(figsize=(9, 6))
    top.plot(kind="barh")
    plt.title(title)
    plt.xlabel("Nombre de commentaires")
    plt.ylabel("Thème")
    plt.tight_layout()
    return fig


# =========================
# Priorisation (sobre)
# =========================
def _normalize(s: str) -> str:
    return ("" if s is None else str(s)).lower().strip()


def classify_category(text: str) -> str:
    t = _normalize(text)
    rules = [
        ("Bug/Crash", ["bug", "crash", "plante", "erreur", "bloque", "impossible", "inutilisable"]),
        ("Performance", ["lent", "lenteur", "chargement", "mouline", "freeze", "fige"]),
        ("UX/Navigation", ["bouton", "écran", "ecran", "onglet", "page", "navigation", "retour", "interface", "menu"]),
        ("Authentification", ["code", "sms", "connexion", "auth", "identifiant", "mot de passe", "login", "face id"]),
        ("Paiement", ["paiement", "payer", "facture", "remboursement", "tarif", "prix", "carte"]),
        ("Documents", ["document", "scan", "justificatif", "pdf"]),
        ("Photos", ["photo", "upload", "télévers", "televers", "ajouter", "pièce jointe", "piece jointe"]),
        ("Support", ["service client", "support", "réponse", "reponse", "appel", "téléphone", "telephone", "mail"]),
        ("Compte", ["compte", "profil", "données", "donnees", "informations", "adresse"]),
    ]
    for cat, kws in rules:
        if any(k in t for k in kws):
            return cat
    return "Autre"


def classify_severity(text: str) -> str:
    t = _normalize(text)
    p0 = ["impossible", "bloque", "bloqué", "crash", "plantage", "erreur", "inutilisable"]
    p1 = ["lent", "lenteur", "bug", "problème", "probleme", "disparu", "perdu", "fige", "freeze"]
    p2 = ["pas clair", "compliqué", "ameliorer", "améliorer", "serait mieux", "dommage"]
    if any(k in t for k in p0): return "P0"
    if any(k in t for k in p1): return "P1"
    if any(k in t for k in p2): return "P2"
    return "P2"


def urgency_score(text: str, sentiment: str = None, csat: float = None) -> int:
    t = _normalize(text)
    score = 10
    strong = ["impossible", "bloque", "crash", "plantage", "erreur", "inutilisable"]
    medium = ["lent", "bug", "problème", "probleme", "disparu", "perdu", "fige", "freeze"]
    if any(k in t for k in strong): score += 50
    if any(k in t for k in medium): score += 25

    s = _normalize(sentiment) if sentiment is not None else ""
    if s in ["neg", "negative", "négatif", "negatif"]: score += 15
    if s in ["pos", "positive", "positif"]: score -= 5

    try:
        if csat is not None and not pd.isna(csat):
            cs = float(csat)
            if cs <= 2: score += 20
            elif cs == 3: score += 8
            elif cs >= 4: score -= 3
    except Exception:
        pass

    return int(max(0, min(100, score)))


def urgency_bucket(score: int) -> str:
    if score >= 75: return "Urgent"
    if score >= 45: return "À traiter"
    return "À surveiller"


def enrich_prioritization(df: pd.DataFrame, text_col: str) -> pd.DataFrame:
    out = df.copy()
    sent_col = "sentiment" if "sentiment" in out.columns else None
    csat_col = "csat" if "csat" in out.columns else None

    out["category"] = out[text_col].astype(str).apply(classify_category)
    out["severity"] = out[text_col].astype(str).apply(classify_severity)

    scores = []
    for i, txt in out[text_col].astype(str).items():
        sent = out.loc[i, sent_col] if sent_col else None
        csat = out.loc[i, csat_col] if csat_col else None
        scores.append(urgency_score(txt, sent, csat))

    out["urgency_score"] = scores
    out["urgency_bucket"] = out["urgency_score"].astype(int).apply(urgency_bucket)
    return out


# =========================
# 3 Ajouts demandés
# =========================
def build_key_points(df_work: pd.DataFrame, text_col: str) -> list[str]:
    n = len(df_work)
    top = df_work["topic_label"].value_counts().head(5)
    main_theme = top.index[0] if len(top) else "—"
    main_count = int(top.iloc[0]) if len(top) else 0
    pct = (main_count / n * 100) if n else 0

    lines = [
        f"Volume analysé : {n} verbatims, {df_work['topic_label'].nunique()} thèmes.",
        f"Irritant principal : « {main_theme} » (~{pct:.0f}% des commentaires).",
        "Décision : prioriser les thèmes récurrents et les P0 (bloquants) dans le backlog.",
    ]

    if "canal" in df_work.columns:
        top_channel = df_work["canal"].astype(str).value_counts().head(1)
        if len(top_channel):
            lines.insert(2, f"Canal le plus touché : {top_channel.index[0]} ({int(top_channel.iloc[0])}).")

    if "sentiment" in df_work.columns:
        neg = (df_work["sentiment"].astype(str).str.lower().isin(["neg", "negative", "négatif", "negatif"])).mean()
        lines.insert(2, f"Part de retours négatifs (approx.) : {neg*100:.0f}%.")

    return lines[:4]


def make_pptx_report(app_name: str, df_work: pd.DataFrame, text_col: str, df_prio: pd.DataFrame | None = None) -> bytes:
    prs = Presentation()

    # Slide 1
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = f"{app_name} — Synthèse"

    box = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(8.4), Inches(4.6))
    tf = box.text_frame
    tf.clear()

    p0 = tf.paragraphs[0]
    p0.text = "Points clés"
    p0.font.bold = True
    p0.font.size = Pt(20)

    for pt in build_key_points(df_work, text_col=text_col):
        p = tf.add_paragraph()
        p.text = f"• {pt}"
        p.font.size = Pt(14)

    # Slide 2
    slide2 = prs.slides.add_slide(prs.slide_layouts[5])
    slide2.shapes.title.text = f"{app_name} — Top thèmes"
    top = df_work["topic_label"].value_counts().head(10)

    box2 = slide2.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(8.4), Inches(4.8))
    tf2 = box2.text_frame
    tf2.clear()
    h = tf2.paragraphs[0]
    h.text = "Top 10 thèmes (volumes)"
    h.font.bold = True
    h.font.size = Pt(18)
    for k, v in top.items():
        p = tf2.add_paragraph()
        p.text = f"• {k} — {int(v)}"
        p.font.size = Pt(14)

    # Slide 3
    slide3 = prs.slides.add_slide(prs.slide_layouts[5])
    slide3.shapes.title.text = f"{app_name} — Backlog priorisé"

    box3 = slide3.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(8.4), Inches(4.8))
    tf3 = box3.text_frame
    tf3.clear()
    h3 = tf3.paragraphs[0]
    h3.text = "Top éléments (à traiter)"
    h3.font.bold = True
    h3.font.size = Pt(18)

    if df_prio is None or "urgency_score" not in df_prio.columns:
        p = tf3.add_paragraph()
        p.text = "• Générer la priorisation (onglet Backlog) pour remplir cette slide."
        p.font.size = Pt(14)
    else:
        top_items = df_prio.sort_values("urgency_score", ascending=False).head(8)
        for _, r in top_items.iterrows():
            sev = r.get("severity", "")
            cat = r.get("category", "")
            score = r.get("urgency_score", "")
            txt = str(r.get(text_col, ""))[:90].replace("\n", " ")
            p = tf3.add_paragraph()
            p.text = f"• {sev} — {cat} — score {score} — {txt}"
            p.font.size = Pt(12)

    bio = io.BytesIO()
    prs.save(bio)
    return bio.getvalue()


def make_pdf_report(df_work: pd.DataFrame, text_col: str, df_prio: pd.DataFrame | None = None) -> bytes:
    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=A4)
    width, height = A4

    c.setFillColorRGB(0.10, 0.20, 0.45)
    c.rect(0, height-3.0*cm, width, 3.0*cm, fill=1, stroke=0)
    c.setFillColorRGB(1, 1, 1)
    c.setFont("Helvetica-Bold", 18)
    c.drawString(2*cm, height-1.6*cm, "INSIGHTIA — Rapport")
    c.setFont("Helvetica", 10)
    c.drawString(2*cm, height-2.2*cm, f"Généré le {datetime.now().strftime('%d/%m/%Y %H:%M')}")

    y = height - 4.0*cm
    c.setFillColorRGB(0, 0, 0)
    c.setFont("Helvetica-Bold", 12)
    c.drawString(2*cm, y, "Points clés")
    y -= 0.6*cm

    c.setFont("Helvetica", 10)
    for line in build_key_points(df_work, text_col=text_col):
        c.drawString(2.2*cm, y, f"• {line[:120]}")
        y -= 0.45*cm

    y -= 0.2*cm
    c.setFont("Helvetica-Bold", 12)
    c.drawString(2*cm, y, "Top thèmes")
    y -= 0.6*cm

    top = df_work["topic_label"].value_counts().head(8)
    c.setFont("Helvetica", 10)
    for t, v in top.items():
        c.drawString(2.2*cm, y, f"• {t[:80]} — {int(v)}")
        y -= 0.45*cm

    y -= 0.2*cm
    c.setFont("Helvetica-Bold", 12)
    c.drawString(2*cm, y, "Backlog (si disponible)")
    y -= 0.6*cm

    c.setFont("Helvetica", 10)
    if df_prio is None:
        c.drawString(2.2*cm, y, "• Générer la priorisation dans l’onglet Backlog pour l’inclure.")
        y -= 0.45*cm
    else:
        p0 = df_prio[df_prio["severity"] == "P0"].sort_values("urgency_score", ascending=False).head(5)
        if p0.empty:
            c.drawString(2.2*cm, y, "• Aucun P0 détecté.")
            y -= 0.45*cm
        else:
            for _, r in p0.iterrows():
                txt = str(r.get(text_col, ""))[:80].replace("\n", " ")
                c.drawString(2.2*cm, y, f"• P0 — {r.get('category','')} — score {r.get('urgency_score','')} — {txt}")
                y -= 0.45*cm

    c.showPage()
    c.save()
    pdf = buffer.getvalue()
    buffer.close()
    return pdf


# =========================
# App layout
# =========================
render_hero()

tabs = st.tabs(["Accueil", "Analyse", "Backlog", "Avant / Après", "Cas d’usage", "Livrables", "Méthode"])

# Session
if "analysis_df" not in st.session_state:
    st.session_state["analysis_df"] = None
    st.session_state["analysis_key"] = None
    st.session_state["analysis_text_col"] = None
    st.session_state["prio_df"] = None


# =========================
# Accueil (storytelling)
# =========================
with tabs[0]:
    c1, c2 = st.columns([1.2, 1])
    with c1:
        card("Ce que fait l’application", "Une lecture claire de la Voix du Client")
        st.markdown("""
- Regroupe automatiquement les verbatims en **thèmes compréhensibles**.  
- Met en évidence les **irritants dominants** avec **exemples**.  
- Construit un **backlog priorisé** (catégorie, sévérité, urgence).  
- Produit des **livrables** partageables (CSV, PDF, PowerPoint).  
        """)
        end_card()

        card("Comment l’utiliser", "Un parcours en 4 étapes")
        st.markdown("""
1) Importer un CSV et choisir la colonne texte  
2) Lancer l’analyse pour obtenir les thèmes  
3) Générer le backlog pour prioriser les actions  
4) Exporter la synthèse pour comité / réunion  
        """)
        end_card()

    with c2:
        card("Démarrer", "Tout est piloté depuis la barre latérale")
        st.markdown("""
Dans l’onglet **Analyse** :
- charge ton fichier CSV
- sélectionne la colonne texte
- lance l’analyse
        """)
        st.markdown("<div class='muted'>Astuce : utilise l’onglet Avant / Après pour mesurer l’impact d’une release.</div>", unsafe_allow_html=True)
        end_card()


# =========================
# Analyse
# =========================
with tabs[1]:
    # Sidebar (propre, stable)
    st.sidebar.markdown("### Import")
    uploaded = st.sidebar.file_uploader("Fichier CSV", type=["csv"])
    if uploaded is None:
        card("Analyse", "Charge un CSV pour démarrer.")
        st.info("Aucun fichier chargé.")
        end_card()
        st.stop()

    auto_sep = detect_sep(uploaded)
    sep = st.sidebar.selectbox("Séparateur", [",", ";", "\t"], index=[",", ";", "\t"].index(auto_sep) if auto_sep in [",", ";", "\t"] else 1)

    try:
        df = read_csv_safely(uploaded, sep=sep)
    except Exception as e:
        card("Erreur", "Lecture CSV impossible")
        st.error(str(e))
        end_card()
        st.stop()

    if df.empty:
        card("Erreur", "Le fichier est vide")
        st.warning("CSV vide.")
        end_card()
        st.stop()

    st.sidebar.markdown("### Paramètres")
    text_col = st.sidebar.selectbox("Colonne texte", options=df.columns.tolist())
    max_rows = st.sidebar.slider("Max lignes analysées", 200, 20000, 2000, 200)
    nr_topics = st.sidebar.slider("Nombre de thèmes", 6, 30, 12, 1)
    sample = st.sidebar.checkbox("Échantillonner si trop grand", value=True)
    run = st.sidebar.button("Lancer l’analyse")

    # Preview section
    card("Données", "Aperçu (contrôle rapide)")
    st.dataframe(df.head(25), use_container_width=True)
    end_card()

    file_bytes = uploaded.getvalue()
    max_rows_effective = min(int(max_rows), 5000)
    key = build_cache_key(file_bytes, sep, text_col, max_rows_effective, int(nr_topics))

    # Cache logic
    if st.session_state["analysis_df"] is not None and st.session_state["analysis_key"] == key and not run:
        df_work = st.session_state["analysis_df"]
    else:
        if not run:
            st.stop()

        df_work = df.copy()
        if len(df_work) > max_rows_effective:
            df_work = df_work.sample(n=max_rows_effective, random_state=42).reset_index(drop=True) if sample else df_work.head(max_rows_effective).reset_index(drop=True)

        with st.spinner("Analyse en cours..."):
            df_work["commentaire_clean"] = clean_text_fast(df_work[text_col].fillna("").astype(str))
            df_work = df_work[df_work["commentaire_clean"].str.strip().ne("")].reset_index(drop=True)

            vectorizer = TfidfVectorizer(max_features=6000, ngram_range=(1, 2))
            X = vectorizer.fit_transform(df_work["commentaire_clean"])

            km = KMeans(n_clusters=int(nr_topics), random_state=42, n_init=10)
            topic_np = km.fit_predict(X)  # numpy
            df_work["topic"] = topic_np

            terms = np.array(vectorizer.get_feature_names_out())
            labels = compute_topic_labels_from_X(X, topic_np, terms, top_n_words=3)
            df_work["topic_label"] = df_work["topic"].map(labels)

        st.session_state["analysis_df"] = df_work
        st.session_state["analysis_key"] = key
        st.session_state["analysis_text_col"] = text_col
        st.session_state["prio_df"] = None

    # Points clés (Ajout 1)
    card("Points clés", "Résumé immédiat pour décideur")
    for line in build_key_points(df_work, text_col=text_col):
        st.write("•", line)
    end_card()

    # KPI section
    st.markdown("## Résultats")
    a, b, c = st.columns(3)
    with a:
        card("Commentaires", "Volume analysé")
        st.metric("", fmt_int(len(df_work)))
        end_card()
    with b:
        card("Thèmes", "Regroupements détectés")
        st.metric("", int(df_work["topic_label"].nunique()))
        end_card()
    with c:
        card("CSAT moyen", "Si disponible")
        if "csat" in df_work.columns:
            csat_mean = pd.to_numeric(df_work["csat"], errors="coerce").mean()
            st.metric("", f"{csat_mean:.2f}" if not np.isnan(csat_mean) else "—")
        else:
            st.metric("", "—")
        end_card()

    # Visuals + detail
    left, right = st.columns([1.2, 1])
    with left:
        card("Thèmes dominants", "Top 15")
        fig = plot_top_topics(df_work["topic_label"], "Top 15 des thèmes", top_n=15)
        st.pyplot(fig, clear_figure=True)
        end_card()

    with right:
        card("Table des thèmes", "Comptage")
        counts = df_work["topic_label"].value_counts().reset_index()
        counts.columns = ["theme", "count"]
        st.dataframe(counts, use_container_width=True, height=420)
        end_card()

    # Evidence section
    card("Exemples", "Preuves terrain")
    top_labels = df_work["topic_label"].value_counts().head(10).index.tolist()
    pick = st.selectbox("Choisir un thème", options=top_labels)
    examples = df_work.loc[df_work["topic_label"] == pick, text_col].dropna().astype(str).head(8).tolist()
    for i, ex in enumerate(examples, 1):
        st.write(f"{i}. {ex}")
    end_card()

    # Export themes
    card("Export", "CSV enrichi avec thème")
    st.download_button(
        "Télécharger le CSV (thèmes)",
        data=df_work.to_csv(index=False).encode("utf-8"),
        file_name="insightia_themes.csv",
        mime="text/csv"
    )
    end_card()


# =========================
# Backlog
# =========================
with tabs[2]:
    card("Backlog", "Priorisation actionnable")
    st.markdown("<div class='muted'>Catégorie, sévérité (P0/P1/P2), urgence.</div>", unsafe_allow_html=True)
    end_card()

    if st.session_state["analysis_df"] is None:
        st.info("Lance d’abord une analyse (onglet Analyse).")
        st.stop()

    df_work = st.session_state["analysis_df"].copy()
    text_col = st.session_state["analysis_text_col"] or df_work.columns[0]

    col1, col2 = st.columns([1, 2])
    with col1:
        generate = st.button("Générer la priorisation")
    with col2:
        st.write("")

    if generate or st.session_state["prio_df"] is not None:
        if st.session_state["prio_df"] is None:
            with st.spinner("Construction du backlog..."):
                st.session_state["prio_df"] = enrich_prioritization(df_work, text_col=text_col)

        df_prio = st.session_state["prio_df"]

        card("Filtres", "Affiner la liste")
        f1, f2, f3 = st.columns(3)
        with f1:
            sev = st.multiselect("Sévérité", ["P0", "P1", "P2"], default=["P0", "P1", "P2"])
        with f2:
            cats = sorted(df_prio["category"].unique().tolist())
            cat_sel = st.multiselect("Catégorie", cats, default=cats)
        with f3:
            urg = st.multiselect("Urgence", ["Urgent", "À traiter", "À surveiller"], default=["Urgent", "À traiter"])
        end_card()

        df_filtered = df_prio[
            (df_prio["severity"].isin(sev)) &
            (df_prio["category"].isin(cat_sel)) &
            (df_prio["urgency_bucket"].isin(urg))
        ].sort_values("urgency_score", ascending=False)

        k1, k2, k3, k4 = st.columns(4)
        with k1:
            card("Total", "")
            st.metric("", fmt_int(len(df_prio)))
            end_card()
        with k2:
            card("Urgent", "")
            st.metric("", int((df_prio["urgency_bucket"] == "Urgent").sum()))
            end_card()
        with k3:
            card("P0", "")
            st.metric("", int((df_prio["severity"] == "P0").sum()))
            end_card()
        with k4:
            card("Catégories", "")
            st.metric("", int(df_prio["category"].nunique()))
            end_card()

        card("Liste priorisée", "Triée par score d’urgence")
        cols_show = [text_col, "topic_label", "category", "severity", "urgency_bucket", "urgency_score"]
        cols_show = [c for c in cols_show if c in df_filtered.columns]
        st.dataframe(df_filtered[cols_show].head(300), use_container_width=True, height=560)
        end_card()

        card("Export", "Backlog (CSV)")
        st.download_button(
            "Télécharger le backlog",
            data=df_prio[cols_show].to_csv(index=False).encode("utf-8"),
            file_name="insightia_backlog.csv",
            mime="text/csv"
        )
        end_card()
    else:
        st.info("Clique sur « Générer la priorisation ».")


# =========================
# Avant / Après
# =========================
with tabs[3]:
    card("Avant / Après", "Comparer deux périodes (impact d’une release)")
    st.markdown("<div class='muted'>Objectif : voir ce qui augmente / diminue entre A et B.</div>", unsafe_allow_html=True)
    end_card()

    left, right = st.columns(2)
    with left:
        card("Période A", "")
        up_a = st.file_uploader("CSV A", type=["csv"], key="csv_a")
        end_card()
    with right:
        card("Période B", "")
        up_b = st.file_uploader("CSV B", type=["csv"], key="csv_b")
        end_card()

    if not up_a or not up_b:
        st.info("Charge les deux fichiers pour activer la comparaison.")
        st.stop()

    sep_a = detect_sep(up_a)
    sep_b = detect_sep(up_b)
    df_a = read_csv_safely(up_a, sep=sep_a)
    df_b = read_csv_safely(up_b, sep=sep_b)

    common_cols = sorted(list(set(df_a.columns).intersection(set(df_b.columns))))
    if not common_cols:
        st.error("Aucune colonne en commun entre les deux fichiers.")
        st.stop()

    card("Paramètres", "")
    text_col_ab = st.selectbox("Colonne texte (commune)", options=common_cols, index=0)
    nr_topics_ab = st.slider("Nombre de thèmes", 6, 30, 12, 1)
    max_rows_ab = st.slider("Max lignes par fichier", 200, 20000, 2000, 200)
    run_ab = st.button("Comparer")
    end_card()

    if not run_ab:
        st.stop()

    def run_express(df_in: pd.DataFrame, text_col: str, nr_topics: int, max_rows: int) -> pd.DataFrame:
        dfx = df_in.copy()
        if len(dfx) > max_rows:
            dfx = dfx.sample(n=max_rows, random_state=42).reset_index(drop=True)

        dfx["commentaire_clean"] = clean_text_fast(dfx[text_col].fillna("").astype(str))
        dfx = dfx[dfx["commentaire_clean"].str.strip().ne("")].reset_index(drop=True)

        vectorizer = TfidfVectorizer(max_features=6000, ngram_range=(1, 2))
        X = vectorizer.fit_transform(dfx["commentaire_clean"])

        km = KMeans(n_clusters=int(nr_topics), random_state=42, n_init=10)
        topic_np = km.fit_predict(X)
        dfx["topic"] = topic_np

        terms = np.array(vectorizer.get_feature_names_out())
        labels = compute_topic_labels_from_X(X, topic_np, terms, top_n_words=3)
        dfx["topic_label"] = dfx["topic"].map(labels)
        return dfx

    with st.spinner("Analyse des deux périodes..."):
        a = run_express(df_a, text_col=text_col_ab, nr_topics=nr_topics_ab, max_rows=max_rows_ab)
        b = run_express(df_b, text_col=text_col_ab, nr_topics=nr_topics_ab, max_rows=max_rows_ab)

    card("KPIs comparatifs", "")
    c1, c2, c3, c4 = st.columns(4)
    with c1: st.metric("A — Commentaires", fmt_int(len(a)))
    with c2: st.metric("B — Commentaires", fmt_int(len(b)))
    with c3: st.metric("A — Thèmes", int(a["topic_label"].nunique()))
    with c4: st.metric("B — Thèmes", int(b["topic_label"].nunique()))
    end_card()

    topA = a["topic_label"].value_counts().head(12)
    topB = b["topic_label"].value_counts().head(12)
    all_labels = sorted(set(topA.index).union(set(topB.index)))

    comp = pd.DataFrame({
        "theme": all_labels,
        "A_count": [int(topA.get(t, 0)) for t in all_labels],
        "B_count": [int(topB.get(t, 0)) for t in all_labels],
    })
    comp["delta"] = comp["B_count"] - comp["A_count"]
    comp = comp.sort_values("delta", ascending=False)

    card("Évolution des thèmes", "Delta = B - A")
    st.dataframe(comp, use_container_width=True, height=460)
    end_card()

    card("Export", "Comparaison (CSV)")
    st.download_button(
        "Télécharger la comparaison",
        data=comp.to_csv(index=False).encode("utf-8"),
        file_name="insightia_avant_apres.csv",
        mime="text/csv"
    )
    end_card()


# =========================
# Cas d’usage (Ajout 2)
# =========================
with tabs[4]:
    card("Cas d’usage", "Où l’application crée de la valeur")
    st.markdown("""
### Produit / UX
- Identifier les irritants dominants et leurs preuves (exemples).
- Prioriser une roadmap correctifs : P0 (bloquants), P1 (gênants), P2 (améliorations).
- Mesurer l’effet d’une release via Avant / Après.

### Support / Opérations
- Repérer les sujets récurrents par canal et réduire les contacts.
- Alimenter une base de connaissance (FAQ) sur les thèmes les plus fréquents.
- Accélérer le traitement grâce à une priorisation objective.

### Direction / Pilotage
- Disposer d’une synthèse : volume, thèmes, urgence, décisions.
- Alimenter un comité avec un livrable (PDF / PowerPoint).
- Objectiver l’évolution des irritants (Avant / Après) pour décider.
    """)
    end_card()


# =========================
# Livrables (PDF + PPTX) (Ajout 3)
# =========================
with tabs[5]:
    card("Livrables", "Formats partageables en réunion")
    if st.session_state["analysis_df"] is None:
        st.info("Lance d’abord une analyse (onglet Analyse) pour générer des livrables.")
        end_card()
        st.stop()

    df_work = st.session_state["analysis_df"]
    text_col = st.session_state["analysis_text_col"] or df_work.columns[0]
    df_prio = st.session_state.get("prio_df", None)

    st.markdown("### Rapport PDF (1 page)")
    if REPORTLAB_AVAILABLE:
        pdf_bytes = make_pdf_report(df_work, text_col=text_col, df_prio=df_prio)
        st.download_button(
            "Télécharger le PDF",
            data=pdf_bytes,
            file_name="insightia_rapport.pdf",
            mime="application/pdf"
        )
    else:
        st.info("PDF indisponible : installe reportlab (pip install reportlab).")

    st.markdown("### PowerPoint (synthèse)")
    if PPTX_AVAILABLE:
        pptx_bytes = make_pptx_report(APP_NAME, df_work=df_work, text_col=text_col, df_prio=df_prio)
        st.download_button(
            "Télécharger le PowerPoint",
            data=pptx_bytes,
            file_name="insightia_synthese.pptx",
            mime="application/vnd.openxmlformats-officedocument.presentationml.presentation"
        )
    else:
        st.info("PowerPoint indisponible : installe python-pptx (pip install python-pptx).")
    end_card()


# =========================
# Méthode
# =========================
with tabs[6]:
    card("Méthode", "Approche robuste et explicable")
    st.markdown("""
- Nettoyage texte léger  
- Vectorisation TF-IDF (unigrams + bigrams)  
- Clustering KMeans (thèmes)  
- Libellés des thèmes par top mots (explicable)  
- Priorisation par règles métier + signaux (sentiment/csat si disponibles)  
    """)
    end_card()


# =========================
# Footer (CDI / sobre)
# =========================
st.markdown("""
<div class="footer">
  <div style="max-width:1100px;margin:auto;">
    <div style="font-weight:950;font-size:16px;">INSIGHTIA — Projet démonstrateur Data & BI</div>
    <div style="color:#64748B;margin-top:6px;font-size:14px;">
      Analyse de verbatims clients → thèmes, priorisation, backlog actionnable (Produit / Support / UX).
    </div>
    <div style="margin-top:14px;font-size:14px;">
      Candidature CDI — Business Analyst (Data & BI)
    </div>
    <div style="margin-top:12px;font-weight:900;">
      AIME Malaka • aimeemalaka84@gmail.com
    </div>
  </div>
</div>
""", unsafe_allow_html=True)
